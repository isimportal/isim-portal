from flask import Flask, request, redirect, session, url_for, send_file
from werkzeug.utils import secure_filename
from flask_mail import Mail, Message
from datetime import datetime, timedelta
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os
import random
import secrets

app = Flask(__name__)
app.secret_key = "gizli-anahtar"

app.config["UPLOAD_FOLDER"] = "static/uploads"
os.makedirs(app.config["UPLOAD_FOLDER"], exist_ok=True)

app.config["MAIL_SERVER"] = "smtp.gmail.com"
app.config["MAIL_PORT"] = 587
app.config["MAIL_USE_TLS"] = True
app.config["MAIL_USERNAME"] = os.environ.get("MAIL_USERNAME")
app.config["MAIL_PASSWORD"] = os.environ.get("MAIL_PASSWORD")
app.config["MAIL_DEFAULT_SENDER"] = ("İŞİM Portal", "isimportal@gmail.com")

mail = Mail(app)

UPDATE_INTERVAL_DAYS = 180

users = {
    "isimportal@gmail.com": {
        "password": "Isimportal2026",
        "role": "admin",
        "verified": True
    }
}

company_data = {}
pending_users = {}
reset_tokens = {}
admin_notifications = []


def send_email(to, subject, body):
    msg = Message(subject=subject, recipients=[to], body=body)
    mail.send(msg)


def is_empty(value):
    return value is None or str(value).strip() == ""


def navbar(menu=""):
    notification_link = ""

    if session.get("role") == "admin":
        unread_count = len([n for n in admin_notifications if not n.get("read", False)])
        notification_link = f'<a href="/notifications">Bildirimler ({unread_count})</a>'

    return f"""
    <div class="navbar">
        <b>İŞİM PORTAL</b>
        <div>
            <a href="/home">Ana Sayfa</a>
            {menu}
            {notification_link}
            <a href="/account">Hesabım</a>
            <a href="/logout">Çıkış</a>
        </div>
    </div>
    """


def page(title, body, menu=""):
    return f"""
    <html>
    <head>
        <title>{title}</title>
        <style>
            body {{
                font-family: Arial;
                background: linear-gradient(135deg, #eef4ff, #ffffff);
                margin: 0;
                color: #222;
            }}

            .navbar {{
                background: linear-gradient(90deg, #1f3f99, #0c2d86);
                color: white;
                padding: 18px 45px;
                display: flex;
                justify-content: space-between;
                align-items: center;
                box-shadow: 0 4px 18px rgba(0,0,0,0.12);
            }}

            .navbar a {{
                color: white;
                text-decoration: none;
                margin-left: 22px;
                font-weight: bold;
                font-size: 15px;
            }}

            .navbar b {{
                font-size: 21px;
                letter-spacing: 1px;
            }}

            .container {{
                width: 90%;
                max-width: 1100px;
                margin: 30px auto;
            }}

            .card {{
                background: white;
                padding: 35px;
                border-radius: 18px;
                box-shadow: 0 4px 18px #ccc;
            }}

            .center-card {{
                background: white;
                width: 68%;
                max-width: 780px;
                margin: 60px auto;
                padding: 46px 40px;
                border-radius: 26px;
                border: 2px solid #1f3f99;
                box-shadow: 0 18px 45px rgba(31,63,153,0.20);
                text-align: center;
            }}

            h1, h2, h3 {{
                color: #003366;
            }}

            h2 {{
                border-bottom: 2px solid #e5e7eb;
                padding-bottom: 8px;
                margin-top: 30px;
            }}

            input, textarea {{
                width: 100%;
                padding: 11px;
                border: 1px solid #ccc;
                border-radius: 8px;
                font-size: 15px;
                box-sizing: border-box;
                margin-top: 6px;
            }}

            textarea {{
                min-height: 90px;
            }}

            label {{
                font-weight: bold;
                display: block;
                margin-bottom: 6px;
            }}

            button, .btn {{
                display: inline-block;
                border: 0;
                border-radius: 8px;
                padding: 12px 18px;
                cursor: pointer;
                font-weight: bold;
                background: #003366;
                color: white;
                text-decoration: none;
                margin-right: 10px;
                margin-bottom: 10px;
            }}

            .btn-edit {{ background: #198754; }}
            .btn-danger {{ background: #dc3545; }}
            .btn-success {{ background: #198754; }}
            .btn-secondary {{ background: #e5e7eb; color: #111; }}

            .grid {{
                display: grid;
                grid-template-columns: repeat(2, 1fr);
                gap: 18px;
            }}

            .full {{
                grid-column: 1 / 3;
            }}

            .section {{
                background: #f8fafc;
                padding: 20px;
                border-radius: 14px;
                margin-bottom: 20px;
                border: 1px solid #e5e7eb;
            }}

            .info-grid {{
                display: grid;
                grid-template-columns: repeat(2, 1fr);
                gap: 15px;
            }}

            .info-item {{
                background: white;
                padding: 14px;
                border-radius: 10px;
                border: 1px solid #e5e7eb;
                box-shadow: 0 2px 6px rgba(0,0,0,0.04);
            }}

            .product-row {{
                display: grid;
                grid-template-columns: 1fr 1fr auto;
                gap: 12px;
                margin-bottom: 12px;
                align-items: end;
            }}

            .company {{
                background: #f8fafc;
                padding: 18px;
                border-radius: 12px;
                margin-bottom: 15px;
                border: 1px solid #e5e7eb;
            }}

            .company:hover {{
                background: #eaf3ff;
            }}

            .company-link {{
                text-decoration: none;
                color: #222;
            }}

            .company-header {{
                display: flex;
                align-items: center;
                gap: 18px;
            }}

            .company-list-logo {{
                width: 90px;
                height: 90px;
                object-fit: contain;
                border-radius: 12px;
                border: 1px solid #ddd;
                background: white;
                padding: 8px;
            }}

            .empty-logo {{
                display: flex;
                align-items: center;
                justify-content: center;
                color: #777;
                font-size: 13px;
                box-sizing: border-box;
            }}

            .company-name {{
                font-size: 24px;
                font-weight: bold;
                color: #003366;
            }}

            .company-logo {{
                max-width: 220px;
                max-height: 140px;
                object-fit: contain;
                border-radius: 12px;
                border: 1px solid #e5e7eb;
                background: white;
                padding: 10px;
            }}

            .logo-area {{
                text-align: center;
                margin: 20px 0;
            }}

            .notification {{
                background: #fff8e1;
                border: 1px solid #f6b51d;
                padding: 20px;
                border-radius: 14px;
                margin-bottom: 18px;
            }}

            .success {{
                background: #d1e7dd;
                color: #0f5132;
                padding: 12px;
                border-radius: 8px;
                margin-bottom: 15px;
            }}

            .error {{
                background: #f8d7da;
                color: #842029;
                padding: 12px;
                border-radius: 8px;
                margin-bottom: 15px;
            }}

            form.inline {{
                display: inline;
            }}
        </style>
    </head>
    <body>
        {navbar(menu)}
        {body}
    </body>
    </html>
    """


def get_company_recipients(company_email, d):
    recipients = []

    if company_email:
        recipients.append(company_email)

    firma_email = d.get("firma_email", "").strip()
    yetkili_email = d.get("yetkili_email", "").strip()

    if firma_email and firma_email not in recipients:
        recipients.append(firma_email)

    if yetkili_email and yetkili_email not in recipients:
        recipients.append(yetkili_email)

    return recipients


def find_missing_fields(d):
    fields = {
        "Firma Adı": d.get("firma_adi", ""),
        "Firma E-posta": d.get("firma_email", ""),
        "Telefon": d.get("telefon", ""),
        "Web Sitesi": d.get("web_sitesi", ""),
        "Adres": d.get("adres", ""),
        "Vergi No": d.get("vergi_no", ""),
        "Üyelik Başlangıç Tarihi": d.get("uyelik_baslangic", ""),
        "Aidat Borcu": d.get("aidat_borcu", ""),
        "Yetkili Ad Soyad": d.get("yetkili_ad", ""),
        "Yetkili E-posta": d.get("yetkili_email", ""),
        "Yetkili Telefon": d.get("yetkili_telefon", ""),
        "Muhasebe Ad Soyad": d.get("muhasebe_ad", ""),
        "Muhasebe E-posta": d.get("muhasebe_email", ""),
        "Muhasebe Telefon": d.get("muhasebe_telefon", ""),
        "Dış Ticaret Ad Soyad": d.get("dis_ticaret_ad", ""),
        "Dış Ticaret E-posta": d.get("dis_ticaret_email", ""),
        "Dış Ticaret Telefon": d.get("dis_ticaret_telefon", ""),
    }

    missing = [name for name, value in fields.items() if is_empty(value)]

    urunler = d.get("urunler", [])
    gtipler = d.get("gtipler", [])

    if not urunler or all(is_empty(x) for x in urunler):
        missing.append("Ürün Adı")

    if not gtipler or all(is_empty(x) for x in gtipler):
        missing.append("GTIP Kodu")

    return missing


def remove_notifications(company_email, notification_type=None):
    global admin_notifications

    if notification_type:
        admin_notifications = [
            n for n in admin_notifications
            if not (n["company_email"] == company_email and n["type"] == notification_type)
        ]
    else:
        admin_notifications = [
            n for n in admin_notifications
            if n["company_email"] != company_email
        ]


def get_notification(company_email, notification_type):
    for n in admin_notifications:
        if n["company_email"] == company_email and n["type"] == notification_type:
            return n
    return None


def send_missing_info_email(company_email, d, missing_fields):
    recipients = get_company_recipients(company_email, d)

    if not recipients:
        return False

    missing_text = "\n".join([f"- {field}" for field in missing_fields])

    mail_body = f"""Sayın ilgili,

İŞİM Portal üzerindeki firma bilgilerinizde bazı eksik alanlar tespit edilmiştir.

Eksik alanlar:
{missing_text}

Lütfen portala giriş yaparak bu alanları doldurunuz.

Eğer bir sorun olduğunu düşünüyorsanız bizimle iletişime geçebilirsiniz.

İŞİM Portal
"""

    for recipient in recipients:
        send_email(
            recipient,
            "İŞİM Portal Eksik Bilgi Bildirimi",
            mail_body
        )

    return True


def send_update_reminder_email(company_email, d):
    recipients = get_company_recipients(company_email, d)

    if not recipients:
        return False

    company_name = d.get("firma_adi", "Firmanız")

    mail_body = f"""Sayın ilgili,

İŞİM Portal üzerinde kayıtlı {company_name} firma bilgilerinizin güncellenme zamanı gelmiştir.

Lütfen portala giriş yaparak firma bilgilerinizi kontrol ediniz ve güncel değilse düzenleyip tekrar kaydediniz.

Bilgileriniz güncelse de, Firma Bilgileri sayfasından Kaydet butonuna basarak güncelleme işlemini tamamlayabilirsiniz.

Eğer bir sorun olduğunu düşünüyorsanız bizimle iletişime geçebilirsiniz.

İŞİM Portal
"""

    for recipient in recipients:
        send_email(
            recipient,
            "İŞİM Portal Firma Bilgileri Güncelleme Hatırlatması",
            mail_body
        )

    return True


def create_missing_info_notification(company_email, d):
    missing_fields = find_missing_fields(d)

    if not missing_fields:
        remove_notifications(company_email, "missing_info")
        return

    old = get_notification(company_email, "missing_info")

    notification = {
        "type": "missing_info",
        "company_email": company_email,
        "company_name": d.get("firma_adi", "Firma adı girilmemiş"),
        "missing_fields": missing_fields,
        "created_at": old["created_at"] if old else datetime.now(),
        "read": old.get("read", False) if old else False,
        "mail_sent": old.get("mail_sent", False) if old else False,
        "last_mail_sent": old.get("last_mail_sent") if old else None
    }

    remove_notifications(company_email, "missing_info")
    admin_notifications.append(notification)


def check_update_reminders():
    for company_email, d in list(company_data.items()):
        if not d or d.get("saved") != "yes":
            continue

        last_updated = d.get("last_updated")

        if not last_updated:
            d["last_updated"] = datetime.now()
            continue

        if isinstance(last_updated, str):
            try:
                last_updated = datetime.fromisoformat(last_updated)
                d["last_updated"] = last_updated
            except Exception:
                d["last_updated"] = datetime.now()
                continue

        if datetime.now() - last_updated >= timedelta(days=UPDATE_INTERVAL_DAYS):
            if get_notification(company_email, "update_due"):
                continue

            notification = {
                "type": "update_due",
                "company_email": company_email,
                "company_name": d.get("firma_adi", "Firma adı girilmemiş"),
                "missing_fields": [],
                "created_at": datetime.now(),
                "read": False,
                "mail_sent": False,
                "last_mail_sent": None,
                "last_updated": last_updated
            }

            admin_notifications.append(notification)

            try:
                send_update_reminder_email(company_email, d)
                notification["mail_sent"] = True
                notification["last_mail_sent"] = datetime.now()
            except Exception as e:
                print("Güncelleme maili gönderilemedi:", e)


def get_selected_companies():
    return session.get("selected_companies", [])


def set_selected_companies(selected):
    session["selected_companies"] = list(dict.fromkeys(selected))


def company_matches_search(email, d, q):
    company_name = d.get("firma_adi", "")
    urunler = " ".join(d.get("urunler", []))
    gtipler = " ".join(d.get("gtipler", []))
    search_text = f"{email} {company_name} {urunler} {gtipler}".lower()
    return q == "" or q in search_text


def create_excel_for_companies(company_emails):
    wb = Workbook()
    ws = wb.active
    ws.title = "Firma Bilgileri"

    headers = [
        "Firma Hesabı", "Firma Adı", "Firma E-posta", "Telefon", "Web Sitesi",
        "Adres", "Vergi No", "Üyelik Başlangıç", "Aidat Borcu",
        "Yetkili Ad Soyad", "Yetkili E-posta", "Yetkili Telefon",
        "Muhasebe Ad Soyad", "Muhasebe E-posta", "Muhasebe Telefon",
        "Dış Ticaret Ad Soyad", "Dış Ticaret E-posta", "Dış Ticaret Telefon",
        "Ürünler", "GTIP Kodları", "Son Güncelleme"
    ]

    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F3F99")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for email in company_emails:
        d = company_data.get(email)

        if not d:
            continue

        last_updated = d.get("last_updated", "")

        if isinstance(last_updated, datetime):
            last_updated = last_updated.strftime("%d.%m.%Y %H:%M")

        row = [
            email,
            d.get("firma_adi", ""),
            d.get("firma_email", ""),
            d.get("telefon", ""),
            d.get("web_sitesi", ""),
            d.get("adres", ""),
            d.get("vergi_no", ""),
            d.get("uyelik_baslangic", ""),
            d.get("aidat_borcu", ""),
            d.get("yetkili_ad", ""),
            d.get("yetkili_email", ""),
            d.get("yetkili_telefon", ""),
            d.get("muhasebe_ad", ""),
            d.get("muhasebe_email", ""),
            d.get("muhasebe_telefon", ""),
            d.get("dis_ticaret_ad", ""),
            d.get("dis_ticaret_email", ""),
            d.get("dis_ticaret_telefon", ""),
            ", ".join(d.get("urunler", [])),
            ", ".join(d.get("gtipler", [])),
            last_updated
        ]

        ws.append(row)

    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {
        "A": 28, "B": 28, "C": 28, "D": 18, "E": 25,
        "F": 40, "G": 18, "H": 18, "I": 15,
        "J": 24, "K": 28, "L": 18,
        "M": 24, "N": 28, "O": 18,
        "P": 24, "Q": 28, "R": 18,
        "S": 35, "T": 25, "U": 22
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    file_stream = BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    return file_stream


@app.route("/")
def login_page():
    return """
    <html>
    <head>
        <style>
            body {font-family: Arial; background: linear-gradient(135deg, #eef4ff, #ffffff); margin: 0;}
            .box {width: 520px; margin: 55px auto; background: white; padding: 38px 42px; border-radius: 24px; border: 2px solid #1f3f99; box-shadow: 0 18px 45px rgba(31,63,153,0.20);}
            .logo-frame {background: white; border-radius: 18px; padding: 18px; margin-bottom: 22px; text-align: center;}
            .logo-login {max-width: 330px; width: 85%;}
            .line {height: 2px; background: #e5e7eb; margin: 22px 0; position: relative;}
            .line::after {content: ""; width: 110px; height: 5px; background: #f6b51d; border-radius: 20px; position: absolute; left: 50%; top: -2px; transform: translateX(-50%);}
            h2 {text-align: center; color: #1f3f99; font-size: 30px; margin-bottom: 22px;}
            input {width: 100%; padding: 14px; margin: 10px 0; border-radius: 10px; border: 1px solid #cfcfcf; box-sizing: border-box; font-size: 16px;}
            button {width: 100%; padding: 14px; border: 0; border-radius: 10px; background: #003fc9; color: white; font-weight: bold; font-size: 16px; cursor: pointer; margin-top: 8px;}
            button:hover {background: #0033a3;}
            .forgot {text-align: center; margin-top: 15px;}
            .forgot a {color: #1f3f99; font-weight: bold; text-decoration: none;}
            hr {margin: 30px 0; border: none; border-top: 1px solid #ddd;}
        </style>
    </head>
    <body>
        <div class="box">
            <div class="logo-frame">
                <img src="/static/uploads/isim_logo.png" class="logo-login">
            </div>

            <div class="line"></div>

            <h2>Giriş Yap</h2>
            <form method="POST" action="/login">
                <input name="email" placeholder="E-posta">
                <input name="password" type="password" placeholder="Şifre">
                <button>Giriş Yap</button>
            </form>

            <div class="forgot">
                <a href="/forgot_password">Şifremi Unuttum</a>
            </div>

            <hr>

            <h2>Firma Kaydı</h2>
            <form method="POST" action="/register">
                <input name="email" placeholder="Firma e-posta">
                <input name="password" type="password" placeholder="Şifre">
                <button>Kayıt Ol</button>
            </form>
        </div>
    </body>
    </html>
    """


@app.route("/register", methods=["POST"])
def register():
    email = request.form["email"].strip()
    password = request.form["password"]

    if email in users:
        return "Bu e-posta zaten kayıtlı. Lütfen giriş yapınız."

    code = str(random.randint(100000, 999999))

    pending_users[email] = {
        "password": password,
        "role": "company",
        "code": code,
        "created_at": datetime.now()
    }

    send_email(
        email,
        "İŞİM Portal E-posta Doğrulama Kodu",
        f"Merhaba,\n\nİŞİM Portal üyeliğinizi tamamlamak için doğrulama kodunuz: {code}\n\nBu kodu doğrulama ekranına girerek hesabınızı aktif edebilirsiniz.\n\nİŞİM Portal"
    )

    return redirect(f"/verify_code?email={email}")


@app.route("/verify_code", methods=["GET", "POST"])
def verify_code():
    email = request.args.get("email", "")

    if request.method == "POST":
        email = request.form["email"].strip()
        code = request.form["code"].strip()

        if email not in pending_users:
            return "Bekleyen kayıt bulunamadı."

        if pending_users[email]["code"] != code:
            return "Doğrulama kodu yanlış."

        users[email] = {
            "password": pending_users[email]["password"],
            "role": "company",
            "verified": True
        }

        company_data[email] = {}
        pending_users.pop(email)

        return """
        <html>
        <body style="font-family:Arial; background:#eef4ff;">
            <div style="width:520px; margin:80px auto; background:white; padding:35px; border-radius:18px;">
                <h2>E-posta doğrulandı. Hesabınız oluşturuldu.</h2>
                <a href="/">Giriş sayfasına dön</a>
            </div>
        </body>
        </html>
        """

    return f"""
    <html>
    <body style="font-family:Arial; background:#eef4ff;">
        <div style="width:420px; margin:80px auto; background:white; padding:35px; border-radius:18px;">
            <h2>E-posta Doğrulama</h2>
            <p>Doğrulama kodu e-posta adresinize gönderildi.</p>
            <form method="POST">
                <input name="email" value="{email}" placeholder="E-posta" style="width:100%; padding:12px; margin-bottom:12px;">
                <input name="code" placeholder="6 haneli doğrulama kodu" style="width:100%; padding:12px; margin-bottom:12px;">
                <button style="width:100%; padding:12px; background:#003fc9; color:white; border:0; border-radius:8px;">
                    Hesabı Doğrula
                </button>
            </form>
        </div>
    </body>
    </html>
    """


@app.route("/login", methods=["POST"])
def login():
    email = request.form["email"].strip()
    password = request.form["password"]

    if email not in users:
        return "Bu e-posta kayıtlı değil."

    if users[email]["password"] != password:
        return "Şifre yanlış."

    if users[email].get("verified") != True:
        return "E-posta adresiniz doğrulanmamış."

    session["email"] = email
    session["role"] = users[email]["role"]

    return redirect("/home")


@app.route("/forgot_password", methods=["GET", "POST"])
def forgot_password():
    if request.method == "POST":
        email = request.form["email"].strip()

        if email not in users:
            return "Bu e-posta ile kayıtlı kullanıcı bulunamadı."

        token = secrets.token_urlsafe(32)

        reset_tokens[token] = {
            "email": email,
            "created_at": datetime.now()
        }

        reset_link = url_for("reset_password", token=token, _external=True)

        send_email(
            email,
            "İŞİM Portal Şifre Sıfırlama Bağlantısı",
            f"Merhaba,\n\nŞifrenizi yenilemek için aşağıdaki bağlantıya tıklayın:\n\n{reset_link}\n\nBu işlemi siz yapmadıysanız bu e-postayı yok sayabilirsiniz.\n\nİŞİM Portal"
        )

        return """
        <html>
        <body style="font-family:Arial; background:#eef4ff;">
            <div style="width:520px; margin:80px auto; background:white; padding:35px; border-radius:18px;">
                <h2>Şifre sıfırlama bağlantısı e-posta adresinize gönderildi.</h2>
                <a href="/">Giriş sayfasına dön</a>
            </div>
        </body>
        </html>
        """

    return """
    <html>
    <body style="font-family:Arial; background:#eef4ff;">
        <div style="width:420px; margin:80px auto; background:white; padding:35px; border-radius:18px;">
            <h2>Şifremi Unuttum</h2>
            <form method="POST">
                <input name="email" placeholder="E-posta" style="width:100%; padding:12px; margin-bottom:12px;">
                <button style="width:100%; padding:12px; background:#003fc9; color:white; border:0; border-radius:8px;">
                    Şifre Sıfırlama Linki Gönder
                </button>
            </form>
            <br>
            <a href="/">Giriş sayfasına dön</a>
        </div>
    </body>
    </html>
    """


@app.route("/reset_password/<token>", methods=["GET", "POST"])
def reset_password(token):
    token_data = reset_tokens.get(token)

    if not token_data:
        return "Geçersiz veya süresi dolmuş şifre sıfırlama bağlantısı."

    if datetime.now() - token_data["created_at"] > timedelta(hours=1):
        reset_tokens.pop(token)
        return "Şifre sıfırlama bağlantısının süresi dolmuş."

    email = token_data["email"]

    if request.method == "POST":
        users[email]["password"] = request.form["password"]
        reset_tokens.pop(token)

        return """
        <html>
        <body style="font-family:Arial; background:#eef4ff;">
            <div style="width:420px; margin:80px auto; background:white; padding:35px; border-radius:18px;">
                <h2>Şifreniz başarıyla değiştirildi.</h2>
                <a href="/">Giriş sayfasına dön</a>
            </div>
        </body>
        </html>
        """

    return """
    <html>
    <body style="font-family:Arial; background:#eef4ff;">
        <div style="width:420px; margin:80px auto; background:white; padding:35px; border-radius:18px;">
            <h2>Yeni Şifre Belirle</h2>
            <form method="POST">
                <input name="password" type="password" placeholder="Yeni şifre" style="width:100%; padding:12px; margin-bottom:12px;">
                <button style="width:100%; padding:12px; background:#003fc9; color:white; border:0; border-radius:8px;">
                    Şifreyi Güncelle
                </button>
            </form>
        </div>
    </body>
    </html>
    """


@app.route("/logout")
def logout():
    session.clear()
    return redirect("/")


@app.route("/account", methods=["GET", "POST"])
def account():
    if "email" not in session:
        return redirect("/")

    current_email = session["email"]
    message = ""

    if request.method == "POST":
        new_email = request.form.get("email", "").strip()
        current_password = request.form.get("current_password", "")
        new_password = request.form.get("new_password", "")

        if current_password != users[current_email]["password"]:
            message = '<div class="error">Mevcut şifre yanlış.</div>'
        else:
            if new_email != current_email:
                if new_email in users:
                    message = '<div class="error">Bu e-posta zaten kullanılıyor.</div>'
                else:
                    users[new_email] = users.pop(current_email)

                    if current_email in company_data:
                        company_data[new_email] = company_data.pop(current_email)

                    session["email"] = new_email
                    current_email = new_email
                    message = '<div class="success">E-posta başarıyla değiştirildi.</div>'

            if new_password:
                users[current_email]["password"] = new_password
                message = '<div class="success">Hesap bilgileri başarıyla güncellendi.</div>'

            if not message:
                message = '<div class="success">Hesap bilgileri başarıyla güncellendi.</div>'

    menu = '<a href="/admin">Admin Paneli</a>' if session["role"] == "admin" else '<a href="/firma_bilgileri">Firma Bilgileri</a>'

    body = f"""
    <div class="container">
        <div class="card" style="max-width:520px; margin:auto;">
            <h1>Hesabım</h1>
            {message}

            <form method="POST">
                <label>E-posta</label>
                <input name="email" value="{current_email}">

                <label>Mevcut Şifre</label>
                <input name="current_password" type="password" placeholder="Değişiklik için mevcut şifreni yaz">

                <label>Yeni Şifre</label>
                <input name="new_password" type="password" placeholder="Değiştirmek istemiyorsan boş bırak">

                <button>Güncelle</button>
            </form>
        </div>
    </div>
    """

    return page("Hesabım", body, menu)


@app.route("/home")
def home():
    if "email" not in session:
        return redirect("/")

    if session["role"] == "admin":
        menu = '<a href="/admin">Admin Paneli</a>'
        extra_text = ""
    else:
        menu = '<a href="/firma_bilgileri">Firma Bilgileri</a>'
        extra_text = '<p>Lütfen üst menüdeki <b>Firma Bilgileri</b> butonuna basarak bilgilerinizi doldurunuz.</p>'

    body = f"""
    <div class="center-card">
        <div style="background:white; border-radius:20px; padding:20px; margin:0 auto 24px auto; max-width:430px;">
            <img src="/static/uploads/isim_logo.png" style="max-width:360px; width:90%;">
        </div>

        <div style="width:120px; height:5px; background:#f6b51d; border-radius:20px; margin:0 auto 24px auto;"></div>

        <h1>İŞİM Portal'a Hoş Geldiniz</h1>

        <div style="font-size:18px; color:#222; font-weight:bold; margin-top:8px;">
            İş ve İnşaat Makineleri Kümesi Yönetim Portalı
        </div>

        {extra_text}
    </div>
    """

    return page("Ana Sayfa", body, menu)


def company_form(d):
    urunler = d.get("urunler", [""])
    gtipler = d.get("gtipler", [""])

    rows = ""
    for i in range(max(len(urunler), 1)):
        urun = urunler[i] if i < len(urunler) else ""
        gtip = gtipler[i] if i < len(gtipler) else ""

        rows += f"""
        <div class="product-row">
            <div><label>Ürün Adı</label><input name="urun_adi" value="{urun}"></div>
            <div><label>GTIP Kodu</label><input name="gtip_kodu" value="{gtip}"></div>
            <button type="button" class="btn-danger" onclick="removeProduct(this)">Sil</button>
        </div>
        """

    logo_preview = ""
    if d.get("logo_filename"):
        logo_preview = f"""
        <div style="margin-top:10px;">
            <p style="color:#198754; font-weight:bold;">Yüklenen Logo: {d.get("logo_filename")}</p>
            <img src="/static/uploads/{d.get("logo_filename")}" style="max-width:180px; max-height:120px; border:1px solid #ddd; border-radius:10px; padding:8px; background:white;">
        </div>
        """

    return f"""
    <div class="container">
        <div class="card">
            <h1>Firma Bilgilerini Düzenle</h1>

            <form method="POST" enctype="multipart/form-data">
                <h2>Genel Bilgiler</h2>
                <div class="grid">
                    <div><label>Firma Adı</label><input name="firma_adi" value="{d.get('firma_adi','')}"></div>
                    <div><label>Firma Logosu</label><input type="file" name="logo">{logo_preview}</div>
                    <div><label>Firma E-posta</label><input name="firma_email" value="{d.get('firma_email','')}"></div>
                    <div><label>Telefon</label><input name="telefon" value="{d.get('telefon','')}"></div>
                    <div><label>Web Sitesi</label><input name="web_sitesi" value="{d.get('web_sitesi','')}"></div>
                    <div><label>Vergi No</label><input name="vergi_no" value="{d.get('vergi_no','')}"></div>
                    <div><label>Üyelik Başlangıç Tarihi</label><input type="date" name="uyelik_baslangic" value="{d.get('uyelik_baslangic','')}"></div>
                    <div><label>Aidat Borcu</label><input name="aidat_borcu" value="{d.get('aidat_borcu','')}"></div>
                    <div class="full"><label>Adres</label><textarea name="adres">{d.get('adres','')}</textarea></div>
                </div>

                <h2>Yetkili Kişi</h2>
                <div class="grid">
                    <div><label>Ad Soyad</label><input name="yetkili_ad" value="{d.get('yetkili_ad','')}"></div>
                    <div><label>E-posta</label><input name="yetkili_email" value="{d.get('yetkili_email','')}"></div>
                    <div><label>Telefon</label><input name="yetkili_telefon" value="{d.get('yetkili_telefon','')}"></div>
                </div>

                <h2>Muhasebe</h2>
                <div class="grid">
                    <div><label>Ad Soyad</label><input name="muhasebe_ad" value="{d.get('muhasebe_ad','')}"></div>
                    <div><label>E-posta</label><input name="muhasebe_email" value="{d.get('muhasebe_email','')}"></div>
                    <div><label>Telefon</label><input name="muhasebe_telefon" value="{d.get('muhasebe_telefon','')}"></div>
                </div>

                <h2>Dış Ticaret Yetkilisi</h2>
                <div class="grid">
                    <div><label>Ad Soyad</label><input name="dis_ticaret_ad" value="{d.get('dis_ticaret_ad','')}"></div>
                    <div><label>E-posta</label><input name="dis_ticaret_email" value="{d.get('dis_ticaret_email','')}"></div>
                    <div><label>Telefon</label><input name="dis_ticaret_telefon" value="{d.get('dis_ticaret_telefon','')}"></div>
                </div>

                <h2>Ürünler ve GTIP Kodları</h2>
                <div id="urunler">{rows}</div>

                <button type="button" class="btn-success" onclick="addProduct()">+ Ürün Ekle</button>

                <div style="display:flex; gap:12px; margin-top:25px; flex-wrap:wrap;">
                    <button>Kaydet</button>
                    <a href="/firma_bilgileri" class="btn btn-secondary">Vazgeç</a>
                </div>
            </form>
        </div>
    </div>

    <script>
        function addProduct() {{
            let div = document.createElement("div");
            div.className = "product-row";
            div.innerHTML = `
                <div><label>Ürün Adı</label><input name="urun_adi"></div>
                <div><label>GTIP Kodu</label><input name="gtip_kodu"></div>
                <button type="button" class="btn-danger" onclick="removeProduct(this)">Sil</button>
            `;
            document.getElementById("urunler").appendChild(div);
        }}

        function removeProduct(btn) {{
            btn.parentElement.remove();
        }}
    </script>
    """


def company_detail(d, self_view=True, company_email=""):
    logo_html = ""
    if d.get("logo_filename"):
        logo_html = f"""
        <div class="logo-area">
            <img src="/static/uploads/{d.get('logo_filename')}" class="company-logo">
        </div>
        """

    urun_html = ""
    urunler = d.get("urunler", [])
    gtipler = d.get("gtipler", [])

    for i in range(len(urunler)):
        gtip = gtipler[i] if i < len(gtipler) else ""
        urun_html += f"<li>📦 {urunler[i]} - 🔢 GTIP: {gtip}</li>"

    if not urun_html:
        urun_html = "<li>Ürün bilgisi girilmemiş.</li>"

    last_updated_text = ""
    if isinstance(d.get("last_updated"), datetime):
        last_updated_text = f"<p><b>Son Güncelleme:</b> {d.get('last_updated').strftime('%d.%m.%Y %H:%M')}</p>"

    if self_view:
        buttons = """
        <a class="btn btn-edit" href="/firma_bilgileri?edit=1">Düzenle</a>
        <form class="inline" method="POST" action="/firma_bilgileri/delete" onsubmit="return confirm('Tüm firma bilgilerinizi silmek istediğinize emin misiniz?')">
            <button class="btn-danger">Tüm Bilgilerimi Sil</button>
        </form>
        <a class="btn" href="/home">Ana Sayfaya Dön</a>
        """
    else:
        buttons = f"""
        <a class="btn" href="/admin">Geri Dön</a>
        <a class="btn btn-success" href="/admin/company/{company_email}/export">Bu Firmayı Excel İndir</a>
        <form class="inline" method="POST" action="/admin/company/{company_email}/delete" onsubmit="return confirm('Bu firmayı sistemden tamamen kaldırmak istediğinize emin misiniz?')">
            <button class="btn-danger">Firmayı Kaldır</button>
        </form>
        """

    return f"""
    <div class="container">
        <div class="card">
            <h1>{d.get('firma_adi','Firma Bilgileri')}</h1>
            {last_updated_text}
            {logo_html}

            <div class="section">
                <h2>Genel Bilgiler</h2>
                <div class="info-grid">
                    <div class="info-item">📧 <b>E-posta:</b> {d.get('firma_email','')}</div>
                    <div class="info-item">📞 <b>Telefon:</b> {d.get('telefon','')}</div>
                    <div class="info-item">🌐 <b>Web Sitesi:</b> {d.get('web_sitesi','')}</div>
                    <div class="info-item">📍 <b>Adres:</b> {d.get('adres','')}</div>
                    <div class="info-item">🧾 <b>Vergi No:</b> {d.get('vergi_no','')}</div>
                    <div class="info-item">📅 <b>Üyelik Başlangıç:</b> {d.get('uyelik_baslangic','')}</div>
                    <div class="info-item">💰 <b>Aidat Borcu:</b> {d.get('aidat_borcu','')}</div>
                </div>
            </div>

            <div class="section">
                <h2>Yetkili Kişi</h2>
                <div class="info-grid">
                    <div class="info-item">👤 <b>Ad Soyad:</b> {d.get('yetkili_ad','')}</div>
                    <div class="info-item">📧 <b>E-posta:</b> {d.get('yetkili_email','')}</div>
                    <div class="info-item">📞 <b>Telefon:</b> {d.get('yetkili_telefon','')}</div>
                </div>
            </div>

            <div class="section">
                <h2>Muhasebe</h2>
                <div class="info-grid">
                    <div class="info-item">👤 <b>Ad Soyad:</b> {d.get('muhasebe_ad','')}</div>
                    <div class="info-item">📧 <b>E-posta:</b> {d.get('muhasebe_email','')}</div>
                    <div class="info-item">📞 <b>Telefon:</b> {d.get('muhasebe_telefon','')}</div>
                </div>
            </div>

            <div class="section">
                <h2>Dış Ticaret Yetkilisi</h2>
                <div class="info-grid">
                    <div class="info-item">👤 <b>Ad Soyad:</b> {d.get('dis_ticaret_ad','')}</div>
                    <div class="info-item">📧 <b>E-posta:</b> {d.get('dis_ticaret_email','')}</div>
                    <div class="info-item">📞 <b>Telefon:</b> {d.get('dis_ticaret_telefon','')}</div>
                </div>
            </div>

            <div class="section">
                <h2>Ürünler ve GTIP Kodları</h2>
                <ul>{urun_html}</ul>
            </div>

            {buttons}
        </div>
    </div>
    """


@app.route("/firma_bilgileri", methods=["GET", "POST"])
def firma_bilgileri():
    if "email" not in session:
        return redirect("/")

    if session["role"] != "company":
        return redirect("/home")

    email = session["email"]

    if request.method == "POST":
        logo = request.files.get("logo")
        logo_filename = company_data.get(email, {}).get("logo_filename", "")

        if logo and logo.filename:
            logo_filename = secure_filename(logo.filename)
            logo.save(os.path.join(app.config["UPLOAD_FOLDER"], logo_filename))

        company_data[email] = {
            "firma_adi": request.form.get("firma_adi", ""),
            "firma_email": request.form.get("firma_email", ""),
            "telefon": request.form.get("telefon", ""),
            "web_sitesi": request.form.get("web_sitesi", ""),
            "adres": request.form.get("adres", ""),
            "vergi_no": request.form.get("vergi_no", ""),
            "uyelik_baslangic": request.form.get("uyelik_baslangic", ""),
            "aidat_borcu": request.form.get("aidat_borcu", ""),
            "yetkili_ad": request.form.get("yetkili_ad", ""),
            "yetkili_email": request.form.get("yetkili_email", ""),
            "yetkili_telefon": request.form.get("yetkili_telefon", ""),
            "muhasebe_ad": request.form.get("muhasebe_ad", ""),
            "muhasebe_email": request.form.get("muhasebe_email", ""),
            "muhasebe_telefon": request.form.get("muhasebe_telefon", ""),
            "dis_ticaret_ad": request.form.get("dis_ticaret_ad", ""),
            "dis_ticaret_email": request.form.get("dis_ticaret_email", ""),
            "dis_ticaret_telefon": request.form.get("dis_ticaret_telefon", ""),
            "urunler": request.form.getlist("urun_adi"),
            "gtipler": request.form.getlist("gtip_kodu"),
            "logo_filename": logo_filename,
            "saved": "yes",
            "last_updated": datetime.now()
        }

        create_missing_info_notification(email, company_data[email])
        remove_notifications(email, "update_due")

        return redirect("/firma_bilgileri")

    d = company_data.get(email, {})
    edit_mode = request.args.get("edit") == "1"

    menu = '<a href="/firma_bilgileri">Firma Bilgileri</a>'

    if not d or d.get("saved") != "yes" or edit_mode:
        return page("Firma Bilgileri", company_form(d), menu)

    return page("Firma Bilgileri", company_detail(d, self_view=True), menu)


@app.route("/firma_bilgileri/delete", methods=["POST"])
def delete_my_company_info():
    if "email" not in session:
        return redirect("/")

    if session["role"] != "company":
        return redirect("/home")

    email = session["email"]

    company_data[email] = {}
    remove_notifications(email)

    return redirect("/firma_bilgileri")


@app.route("/notifications")
def notifications():
    if "email" not in session or session["role"] != "admin":
        return redirect("/home")

    check_update_reminders()

    for n in admin_notifications:
        n["read"] = True

    notification_html = ""

    if not admin_notifications:
        notification_html = "<p>Şu anda bildirim bulunmamaktadır.</p>"
    else:
        for index, n in enumerate(reversed(admin_notifications)):
            real_index = len(admin_notifications) - 1 - index

            mail_status = "Gönderilmedi"
            if n.get("mail_sent") and n.get("last_mail_sent"):
                mail_status = f"Gönderildi - {n['last_mail_sent'].strftime('%d.%m.%Y %H:%M')}"

            if n["type"] == "missing_info":
                missing_list = "".join([f"<li>{field}</li>" for field in n["missing_fields"]])
                content = f"""
                <p><b>Bildirim türü:</b> Eksik bilgi</p>
                <p><b>Firma hesabı:</b> {n["company_email"]}</p>
                <p><b>Durum:</b> Hâlâ eksik bilgi var</p>
                <p><b>Mail durumu:</b> {mail_status}</p>
                <p><b>Eksik alanlar:</b></p>
                <ul>{missing_list}</ul>
                """
                button_text = "Eksik Bilgi Maili Gönder"
            else:
                last_updated = n.get("last_updated")
                last_updated_text = last_updated.strftime("%d.%m.%Y %H:%M") if isinstance(last_updated, datetime) else "Bilinmiyor"

                content = f"""
                <p><b>Bildirim türü:</b> 6 Aylık Güncelleme Hatırlatması</p>
                <p><b>Firma hesabı:</b> {n["company_email"]}</p>
                <p><b>Durum:</b> Mail gönderildiği halde firma bilgilerini hâlâ güncellemedi</p>
                <p><b>Son güncelleme:</b> {last_updated_text}</p>
                <p><b>Mail durumu:</b> {mail_status}</p>
                """
                button_text = "Güncelleme Maili Gönder"

            notification_html += f"""
            <div class="notification">
                <h3>{n["company_name"]}</h3>
                {content}
                <small>Bildirim tarihi: {n["created_at"].strftime("%d.%m.%Y %H:%M")}</small>
                <br><br>
                <form method="POST" action="/notifications/send/{real_index}">
                    <button class="btn-success">{button_text}</button>
                </form>
            </div>
            """

    body = f"""
    <div class="container">
        <div class="card">
            <h1>Bildirimler</h1>
            {notification_html}
            <br>
            <a class="btn" href="/admin">Admin Paneline Dön</a>
        </div>
    </div>
    """

    return page("Bildirimler", body, '<a href="/admin">Admin Paneli</a>')


@app.route("/notifications/send/<int:index>", methods=["POST"])
def send_notification_mail(index):
    if "email" not in session or session["role"] != "admin":
        return redirect("/home")

    if index < 0 or index >= len(admin_notifications):
        return redirect("/notifications")

    n = admin_notifications[index]
    company_email = n["company_email"]
    d = company_data.get(company_email, {})

    try:
        if n["type"] == "missing_info":
            missing_fields = find_missing_fields(d)

            if not missing_fields:
                remove_notifications(company_email, "missing_info")
                return redirect("/notifications")

            send_missing_info_email(company_email, d, missing_fields)
            n["missing_fields"] = missing_fields

        elif n["type"] == "update_due":
            send_update_reminder_email(company_email, d)

        n["mail_sent"] = True
        n["last_mail_sent"] = datetime.now()
        n["read"] = False

    except Exception as e:
        print("Bildirim maili gönderilemedi:", e)

    return redirect("/notifications")


@app.route("/admin")
def admin():
    if "email" not in session or session["role"] != "admin":
        return redirect("/home")

    check_update_reminders()

    q = request.args.get("q", "").lower().strip()
    selected_companies = get_selected_companies()

    company_list = ""
    visible_companies = []

    for email, d in company_data.items():
        if company_matches_search(email, d, q):
            visible_companies.append(email)
            company_name = d.get("firma_adi", "Firma adı girilmemiş")

            if d.get("logo_filename"):
                logo_html = f'<img src="/static/uploads/{d.get("logo_filename")}" class="company-list-logo">'
            else:
                logo_html = '<div class="company-list-logo empty-logo">Logo Yok</div>'

            checked = "checked" if email in selected_companies else ""

            company_list += f"""
            <div class="company">
                <div class="company-header">
                    <input type="checkbox" name="selected_visible" value="{email}" {checked}
                           style="width:22px; height:22px; margin-right:5px;">

                    {logo_html}

                    <a class="company-link" href="/admin/company/{email}" style="flex:1;">
                        <div class="company-name">{company_name}</div>
                        <div style="color:#555; margin-top:6px;">{email}</div>
                    </a>
                </div>
            </div>
            """

    if company_list == "":
        company_list = "<p>Aramanıza uygun firma bulunamadı.</p>"

    hidden_visible = ""
    for email in visible_companies:
        hidden_visible += f'<input type="hidden" name="visible_companies" value="{email}">'

    if selected_companies:
        selected_items = ""

        for email in selected_companies:
            d = company_data.get(email, {})
            selected_items += f"""
            <li>
                <b>{d.get("firma_adi", "Firma adı girilmemiş")}</b>
                <span style="color:#555;">({email})</span>
            </li>
            """

        selected_html = f"""
        <div class="section">
            <h2>Seçilen Firmalar</h2>
            <p><b>Seçilen firma sayısı:</b> {len(selected_companies)}</p>
            <ul>{selected_items}</ul>

            <a class="btn btn-danger" href="/admin/clear_selection">
                Seçimleri Temizle
            </a>
        </div>
        """
    else:
        selected_html = """
        <div class="section">
            <h2>Seçilen Firmalar</h2>
            <p>Henüz firma seçilmedi.</p>
        </div>
        """

    body = f"""
    <div class="container">
        <div class="card">
            <h1>Admin Paneli</h1>
            <p>Kayıtlı firmaları şirket adı, ürün adı veya GTIP kodu ile arayabilirsiniz.</p>

            <form method="GET" style="display:flex; gap:10px; margin-bottom:25px;">
                <input name="q" placeholder="Firma adı, ürün adı veya GTIP kodu ara" value="{q}">
                <button>Ara</button>
            </form>

            {selected_html}

            <form method="POST" action="/admin/update_selection">
                <input type="hidden" name="q" value="{q}">
                {hidden_visible}

                <div style="margin-bottom:15px;">
                    <button type="button" class="btn btn-secondary" onclick="selectAllVisible()">
                        Görünenlerin Hepsini Seç
                    </button>

                    <button type="button" class="btn btn-secondary" onclick="clearVisibleSelection()">
                        Görünen Seçimleri Kaldır
                    </button>

                    <button name="action" value="save" class="btn btn-success">
                        Seçimi Kaydet
                    </button>

                    <button name="action" value="export" class="btn btn-success">
                        Seçilenleri Excel İndir
                    </button>
                </div>

                {company_list}
            </form>
        </div>
    </div>

    <script>
        function selectAllVisible() {{
            document.querySelectorAll('input[name="selected_visible"]').forEach(function(cb) {{
                cb.checked = true;
            }});
        }}

        function clearVisibleSelection() {{
            document.querySelectorAll('input[name="selected_visible"]').forEach(function(cb) {{
                cb.checked = false;
            }});
        }}
    </script>
    """

    return page("Admin Paneli", body, '<a href="/admin">Admin Paneli</a>')


@app.route("/admin/update_selection", methods=["POST"])
def update_selection():
    if "email" not in session or session["role"] != "admin":
        return redirect("/home")

    q = request.form.get("q", "")
    action = request.form.get("action", "save")

    visible_companies = request.form.getlist("visible_companies")
    selected_visible = request.form.getlist("selected_visible")

    selected_companies = get_selected_companies()

    selected_companies = [
        email for email in selected_companies
        if email not in visible_companies
    ]

    for email in selected_visible:
        if email not in selected_companies:
            selected_companies.append(email)

    selected_companies = [
        email for email in selected_companies
        if email in company_data
    ]

    set_selected_companies(selected_companies)

    if action == "export":
        if not selected_companies:
            return "Excel oluşturmak için önce firma seçmelisiniz."

        file_stream = create_excel_for_companies(selected_companies)
        filename = f"secilen_firmalar_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

        return send_file(
            file_stream,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    return redirect(f"/admin?q={q}")


@app.route("/admin/clear_selection")
def clear_selection():
    if "email" not in session or session["role"] != "admin":
        return redirect("/home")

    session["selected_companies"] = []

    return redirect("/admin")


@app.route("/admin/export_selected")
def export_selected_companies():
    if "email" not in session or session["role"] != "admin":
        return redirect("/home")

    selected_companies = [
        email for email in get_selected_companies()
        if email in company_data
    ]

    if not selected_companies:
        return "Excel oluşturmak için önce firma seçmelisiniz."

    file_stream = create_excel_for_companies(selected_companies)

    filename = f"secilen_firmalar_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    return send_file(
        file_stream,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.route("/admin/company/<company_email>/export")
def export_single_company(company_email):
    if "email" not in session or session["role"] != "admin":
        return redirect("/home")

    if company_email not in company_data:
        return "Firma bulunamadı."

    file_stream = create_excel_for_companies([company_email])

    filename = f"firma_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    return send_file(
        file_stream,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.route("/admin/company/<company_email>")
def admin_company_detail(company_email):
    if "email" not in session or session["role"] != "admin":
        return redirect("/home")

    check_update_reminders()

    d = company_data.get(company_email)

    if not d:
        return "Firma bulunamadı."

    return page(
        "Firma Detayı",
        company_detail(d, self_view=False, company_email=company_email),
        '<a href="/admin">Admin Paneli</a>'
    )


@app.route("/admin/company/<company_email>/delete", methods=["POST"])
def admin_delete_company(company_email):
    if "email" not in session or session["role"] != "admin":
        return redirect("/home")

    if company_email in company_data:
        company_data.pop(company_email)

    if company_email in users and users[company_email].get("role") == "company":
        users.pop(company_email)

    remove_notifications(company_email)

    selected_companies = get_selected_companies()
    selected_companies = [email for email in selected_companies if email != company_email]
    set_selected_companies(selected_companies)

    return redirect("/admin")


if __name__ == "__main__":
    app.run(debug=True)